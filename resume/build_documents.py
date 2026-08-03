"""Renders the Campminder application markdown into openable PDF and Word files.

Usage: python3 resume/build_documents.py
"""

import re
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)

DOCUMENTS_DIR = Path(__file__).parent
RESUME_STEM = "forest-yox-resume-campminder"
NOTES_STEM = "campminder-positioning-notes"

BODY_FONT = "Times-Roman"
BOLD_FONT = "Times-Bold"


def centered_headings_styles():
    return build_styles(heading_alignment=TA_CENTER)


def left_headings_styles():
    return build_styles(heading_alignment=TA_LEFT)


def build_styles(heading_alignment):
    return {
        "name": ParagraphStyle(
            "name", fontName=BOLD_FONT, fontSize=21, leading=25, alignment=TA_CENTER, spaceAfter=4
        ),
        "contact": ParagraphStyle(
            "contact", fontName=BODY_FONT, fontSize=10.5, leading=13, alignment=TA_CENTER, spaceAfter=6
        ),
        "headline": ParagraphStyle(
            "headline", fontName=BODY_FONT, fontSize=11, leading=14, alignment=TA_CENTER, spaceAfter=7
        ),
        "body": ParagraphStyle(
            "body", fontName=BODY_FONT, fontSize=10, leading=12.8, alignment=TA_LEFT, spaceAfter=5
        ),
        "section": ParagraphStyle(
            "section",
            fontName=BOLD_FONT,
            fontSize=12,
            leading=15,
            alignment=heading_alignment,
            spaceBefore=11,
            spaceAfter=6,
        ),
        "subsection": ParagraphStyle(
            "subsection",
            fontName=BOLD_FONT,
            fontSize=10.5,
            leading=13,
            alignment=TA_LEFT,
            spaceBefore=9,
            spaceAfter=4,
        ),
        "bullet": ParagraphStyle(
            "bullet", fontName=BODY_FONT, fontSize=10, leading=12.6, alignment=TA_LEFT, spaceAfter=2.5
        ),
        "cell": ParagraphStyle(
            "cell", fontName=BODY_FONT, fontSize=8.5, leading=10.8, alignment=TA_LEFT
        ),
        "role": ParagraphStyle(
            "role",
            fontName=BODY_FONT,
            fontSize=10,
            leading=12.8,
            alignment=TA_LEFT,
            spaceBefore=7,
            spaceAfter=4,
        ),
    }


def to_inline_markup(text):
    escaped = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    coded = re.sub(r"`(.+?)`", r'<font face="Courier" size="9">\1</font>', escaped)
    bolded = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", coded)
    return re.sub(r"(?<!\*)\*([^*]+?)\*(?!\*)", r"<i>\1</i>", bolded)


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_divider(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def parse_blocks(markdown_text):
    """Groups hard-wrapped markdown lines into (kind, payload) blocks."""
    blocks = []
    pending_paragraph = []

    def flush_paragraph():
        if pending_paragraph:
            blocks.append(("paragraph", " ".join(pending_paragraph)))
            pending_paragraph.clear()

    for line in markdown_text.splitlines():
        stripped = line.strip()
        continues_bullet = line.startswith("  ") and blocks and blocks[-1][0] in ("bullet", "numbered")
        starts_labelled_line = stripped.startswith("**") and pending_paragraph

        if not stripped or starts_labelled_line or stripped.startswith("|"):
            flush_paragraph()

        if not stripped:
            continue

        numbered = re.match(r"(\d+)\.\s+(.*)", stripped)

        if stripped.startswith("# "):
            blocks.append(("name", stripped[2:]))
        elif stripped.startswith("## "):
            blocks.append(("section", stripped[3:]))
        elif stripped.startswith("### "):
            blocks.append(("subsection", stripped[4:]))
        elif is_table_divider(stripped):
            continue
        elif stripped.startswith("|"):
            kind = "table_row" if blocks and blocks[-1][0] in ("table_row", "table_header") else "table_header"
            blocks.append((kind, split_table_row(stripped)))
        elif stripped.startswith("- "):
            blocks.append(("bullet", stripped[2:]))
        elif numbered:
            blocks.append(("numbered", numbered.group(2)))
        elif continues_bullet and not pending_paragraph:
            blocks[-1] = (blocks[-1][0], f"{blocks[-1][1]} {stripped}")
        else:
            pending_paragraph.append(stripped)

    flush_paragraph()
    return blocks


def to_flowables(blocks, styles, available_width):
    flowables = []
    pending_list = []
    pending_list_kind = None
    pending_table = []
    paragraphs_before_first_section = 0
    reached_first_section = False
    unattached_role_header = None

    def item_list(texts, kind, space_after):
        return ListFlowable(
            [ListItem(Paragraph(text, styles["bullet"]), leftIndent=16) for text in texts],
            bulletType="1" if kind == "numbered" else "bullet",
            bulletFontSize=9 if kind == "numbered" else 6,
            bulletOffsetY=0 if kind == "numbered" else -1.5,
            start="1" if kind == "numbered" else "circle",
            leftIndent=16,
            spaceAfter=space_after,
        )

    def flush_list():
        """Binds a role header to its first item so headers never orphan at a page break."""
        nonlocal unattached_role_header, pending_list_kind
        if not pending_list:
            return
        if unattached_role_header is not None:
            flowables.append(
                KeepTogether([unattached_role_header, item_list(pending_list[:1], pending_list_kind, 2.5)])
            )
            unattached_role_header = None
            remaining = pending_list[1:]
        else:
            remaining = pending_list[:]
        if remaining:
            flowables.append(item_list(remaining, pending_list_kind, 4))
        pending_list.clear()
        pending_list_kind = None

    def flush_table():
        if not pending_table:
            return
        column_width = available_width / len(pending_table[0])
        rendered_rows = [
            [Paragraph(to_inline_markup(cell), styles["cell"]) for cell in row] for row in pending_table
        ]
        table = Table(rendered_rows, colWidths=[column_width] * len(pending_table[0]))
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#999999")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        flowables.append(table)
        pending_table.clear()

    for kind, payload in blocks:
        if kind in ("bullet", "numbered"):
            if pending_list_kind and pending_list_kind != kind:
                flush_list()
            pending_list_kind = kind
            pending_list.append(to_inline_markup(payload))
            continue
        if kind in ("table_header", "table_row"):
            flush_list()
            pending_table.append(payload)
            continue

        flush_list()
        flush_table()
        if unattached_role_header is not None:
            flowables.append(unattached_role_header)
            unattached_role_header = None

        text = to_inline_markup(payload)
        if kind == "name":
            flowables.append(Paragraph(text, styles["name"]))
        elif kind == "section":
            reached_first_section = True
            flowables.append(Paragraph(text, styles["section"]))
        elif kind == "subsection":
            flowables.append(Paragraph(text, styles["subsection"]))
        elif not reached_first_section:
            paragraphs_before_first_section += 1
            style = {1: "contact", 2: "headline"}.get(paragraphs_before_first_section, "body")
            flowables.append(Paragraph(text, styles[style]))
        elif payload.startswith("**") and " | " in payload:
            unattached_role_header = Paragraph(text, styles["role"])
        else:
            flowables.append(Paragraph(text, styles["body"]))

    flush_list()
    flush_table()
    if unattached_role_header is not None:
        flowables.append(unattached_role_header)
    return flowables


def build_pdf(markdown_path, pdf_path, styles, title):
    document = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        leftMargin=0.72 * inch,
        rightMargin=0.72 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.55 * inch,
        title=title,
        author="Forest Yox",
    )
    blocks = parse_blocks(markdown_path.read_text())
    document.build(to_flowables(blocks, styles, document.width))
    print(f"Wrote {pdf_path}")


def add_runs_with_emphasis(paragraph, text):
    for segment in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`.+?`)", text):
        if not segment:
            continue
        run = paragraph.add_run(re.sub(r"^(\*\*|\*|`)|(\*\*|\*|`)$", "", segment))
        run.bold = segment.startswith("**")
        run.italic = segment.startswith("*") and not segment.startswith("**")
        if segment.startswith("`"):
            run.font.name = "Consolas"


def build_docx(markdown_path, docx_path):
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)

    blocks = parse_blocks(markdown_path.read_text())
    paragraphs_before_first_section = 0
    reached_first_section = False

    for kind, payload in blocks:
        if kind in ("table_header", "table_row"):
            continue
        if kind == "name":
            heading = document.add_paragraph()
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = heading.add_run(payload)
            run.bold = True
            run.font.size = Pt(21)
            continue
        if kind == "section":
            reached_first_section = True
            heading = document.add_paragraph()
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = heading.add_run(payload)
            run.bold = True
            run.font.size = Pt(12)
            continue

        paragraph = document.add_paragraph(style="List Bullet" if kind == "bullet" else None)
        if not reached_first_section and kind == "paragraph":
            paragraphs_before_first_section += 1
            if paragraphs_before_first_section <= 2:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs_with_emphasis(paragraph, payload)

    document.save(str(docx_path))
    print(f"Wrote {docx_path}")


if __name__ == "__main__":
    resume_markdown = DOCUMENTS_DIR / f"{RESUME_STEM}.md"
    notes_markdown = DOCUMENTS_DIR / f"{NOTES_STEM}.md"

    build_pdf(resume_markdown, DOCUMENTS_DIR / f"{RESUME_STEM}.pdf", centered_headings_styles(), "Forest Yox — Resume")
    build_docx(resume_markdown, DOCUMENTS_DIR / f"{RESUME_STEM}.docx")
    build_pdf(
        notes_markdown,
        DOCUMENTS_DIR / f"{NOTES_STEM}.pdf",
        left_headings_styles(),
        "Campminder positioning notes",
    )
